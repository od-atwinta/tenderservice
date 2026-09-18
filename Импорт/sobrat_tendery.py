import openpyxl, re, json, datetime, collections, sys, warnings
warnings.filterwarnings('ignore')
SRC = sys.argv[1]; OUT = sys.argv[2]; REPORT = sys.argv[3]
wb = openpyxl.load_workbook(SRC)
ws = wb['Тендеры']; wa = wb['Ответы']
H = [c.value for c in ws[1]]; C = H.index
log = collections.defaultdict(list)

MANAGERS = {'денисенко':'Оксана Денисенко','долинин':'Никита Долинин','величко':'Никита Долинин',
            'мильберг':'Владислав Мильберг','юров':'Дмитрий Юров','щетинина':'Дарья Щетинина',
            'ельцова':'Ксения Ельцова','полковников':'Андрей Полковников'}
DEPTS = {'ТП':'Техническая поддержка'}
REASONS = ['Нейронка плохо определила','Прочее','Не наш стек','Не наш стек/профиль','Не успеваем подготовиться',
           'Не проходим по КТ','Жесткие условия','Малобюджетный','Нереальные сроки','Отмена заказчиком',
           'Закрытая от нас','Дубль','Не успели податься','Невыполнимые условия']
LOSS = ['Отклонено','Отменено','Без выбора победителя','Завершено без ответа','Не допущены','Отказались сами']
FINAL = LOSS + ['Выиграли']

def s(v): return str(v).strip() if v not in (None,'') else ''
def manager(v):
    k = s(v).lower()
    if not k: return ''
    if k not in MANAGERS: log['unknown_manager'].append(k); return s(v)
    return MANAGERS[k]
def reason(v):
    k = s(v)
    if not k: return ''
    for r in REASONS:
        if r.lower() == k.lower(): return r
    log['unknown_reason'].append(k); return k
def dept(v):
    k = s(v)
    return DEPTS.get(k, k)
def iso_date(v, row, field):
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.date): return v.strftime('%Y-%m-%d')
    t = s(v)
    if not t: return ''
    m = re.search(r'(\d{1,2})\.(\d{1,2})\.(\d{2,4})', t)
    if m:
        y = int(m.group(3)); y = y + 2000 if y < 100 else y
        return '%04d-%02d-%02d' % (y, int(m.group(2)), int(m.group(1)))
    log['bad_date'].append((row, field, t)); return ''
def iso_deadline(v, row):
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%dT%H:%M:00') if (v.hour or v.minute) else v.strftime('%Y-%m-%d')
    return iso_date(v, row, 'deadline')
def money(v, row):
    if isinstance(v, (int, float)): return float(v) if v > 0 else None
    t = s(v)
    if not t: return None
    t2 = t.replace('\xa0','').replace(' ','').replace(',','.')
    try:
        n = float(t2); return n if n > 0 else None
    except ValueError:
        log['bad_nmc'].append((row, t)); return None
def tid(u):
    m = re.search(r'tender=([0-9a-f]+)', s(u)); return m.group(1) if m else None
def norm(t): return re.sub(r'\W+', '', s(t).lower())

# ---------- main rows ----------
rows = {}
for i, r in enumerate(ws.iter_rows(min_row=2, max_col=27, values_only=True), start=2):
    if not s(r[1]) and not (s(r[2]) and s(r[C('Решение об участии')])): continue
    if s(r[1]) and not s(r[0]) and all(v is None for v in r[2:22]): continue  # месяц-разделитель
    d = s(r[C('Решение об участии')])
    if not d and not s(r[2]): continue
    rows[i] = r
excluded = {i for i, r in rows.items()
            if s(r[C('Решение об участии')]) == 'дубль'
            or (s(r[C('Решение об участии')]) == 'Отказ' and s(r[C('дубль?')]))}
empty_dec = [i for i, r in rows.items() if not s(r[C('Решение об участии')])]
if empty_dec: log['empty_decision'] = empty_dec

# ---------- answers ----------
MANUAL = {4:2226, 12:2212, 90:1152, 62:12, 48:1210, 36:2246, 22:2399, 59:3, 65:4}
NEW_FROM_ANSWERS = [104]
SUBMIT_DATE_FROM_STATS = {1565:'2026-04-14',  # «статистика», 14.04: «Мерси»
                          1440:'2026-04-15'}  # «статистика», 15.04: «Ламода - аналитика»; в «Ответах» 25.08 — основной этап
byid = collections.defaultdict(set); byname = collections.defaultdict(set)
for i, r in rows.items():
    if i in excluded: continue
    if tid(r[11]): byid[tid(r[11])].add(i)
    byname[norm(r[1])].add(i)
answers = {}
answer_rows = {}
for i, r in enumerate(wa.iter_rows(min_row=2, max_col=10, values_only=True), start=2):
    if not any(v not in (None, '') for v in r): continue
    answer_rows[i] = r
    if i in NEW_FROM_ANSWERS: continue
    if i in MANUAL: m = MANUAL[i]
    else:
        cand = byid.get(tid(r[4])) or byname.get(norm(r[1])) or set()
        if len(cand) != 1: log['answer_unmatched'].append((i, s(r[1])[:60], sorted(cand))); continue
        m = next(iter(cand))
    if m in answers: log['answer_double'].append((i, m, answers[m][0]))
    answers[m] = (i, r)

# ---------- stages from «статистика» ----------
def D(dm): return '2026-%s-%s' % (dm[3:5], dm[0:2])
RETENDER = [(4,'01.02'),(764,'06.03'),(921,'13.03'),(836,'14.04'),(1472,'28.04'),(1442,'30.04'),(1646,'02.06'),
            (2062,'07.07'),(2000,'07.07'),(1847,'10.07'),(2114,'13.07'),(2157,'22.07'),(2200,'31.07'),
            (2096,'05.08'),(2356,'03.09'),(2361,'10.09'),(1847,'10.09')]
MAINSTAGE = [(76,'01.02'),(529,'19.02'),(921,'02.03'),(479,'02.03'),(836,'30.03'),(1440,'25.08'),(2401,'08.09'),(2266,'15.09')]
extra = collections.defaultdict(list)
for row, dm in MAINSTAGE: extra[row].append(('Этап 2', D(dm)))
for row, dm in RETENDER: extra[row].append(('Переторжка', D(dm)))
for row in extra: extra[row].sort(key=lambda x: x[1])

TENDERS_STATUS = {'Подали':'Подано','Формируем':'Формирование заявки','ПКО/НДА':'Подано','Переторжка':'Подано',
    'Отклонена':'Отклонено','не успели податься':'Отказались сами','б/выбора победителя':'Без выбора победителя',
    'Одобрена':'Выиграли','Контракт заключен':'Выиграли','Отменено':'Отменено','Второй этап':'Подано',
    'Отказались сами':'Отказались сами'}
ANSWER_STATUS = {'Отказ':'Отклонено','без выбора':'Без выбора победителя','RFI без ответа':'Завершено без ответа',
    'Выиграли':'Выиграли','Второй этап':'Подано','Основной этап':'Подано'}
PROCHEE = {2000:'Отменено',1038:'Отменено',2103:'Отменено',1210:'Без выбора победителя',1821:'Отклонено',
           1917:'Завершено без ответа',2164:'Завершено без ответа'}
PKO_REFUSED = {1430, 1746, 1841, 2442, 2238}
# дата победы — из примечаний к столбцу «выиграли» листа «статистика»
WIN_DATES = {836:'2026-05-14', 1647:'2026-05-15', 1646:'2026-06-22', 2062:'2026-07-09'}
# этап ПКО/НДА перед основной подачей (примечания «статистики»)
PRE_STAGE = {1442: ('НДА', '2026-04-06')}
# первая подача была на ПКО («статистика»: «Дрогери ритейл - ПКО», «ОЗОН игровые механики, подались на ПКО»)
FIRST_IS_PKO = {1471, 1394}
NOT_SUBMITTED_STATUSES = {'Формирование заявки'}
# отказ заказчика после подачи: решение в «Тендерах» — «Отказ», но подача была («Ответы» №77, 18.02)
SUBMITTED_REFUSALS = {899}

def section(idx, name, status, deadline, actual, completed):
    return {'id': None, 'name': name, 'status': status, 'deadline': deadline,
            'actualSubmissionDate': actual, 'completed': completed}

out = []
nid = 0
for i in sorted(rows):
    if i in excluded: continue
    r = rows[i]; nid += 1
    dec = s(r[C('Решение об участии')]); st_raw = s(r[C('Статус заявки')])
    ans = answers.get(i); ar = ans[1] if ans else None
    res = s(ar[8]) if ar else ''
    t = {'id': nid, 'sheetRow': i,
         'manager': manager(ar[0]) if ar and s(ar[0]) else manager(r[0]),
         'name': s(r[1]) or (s(ar[1]) if ar else ''), 'customer': s(r[2]), 'nmc': money(r[4], i), 'currencies': ['₽'],
         'pub': iso_date(r[9], i, 'pub'), 'deadline': iso_deadline(r[10], i),
         'dept': dept(r[C('Отдел')]), 'depts': [dept(r[C('Отдел')])] if s(r[C('Отдел')]) else [],
         'decision': dec, 'rejectReason': reason(r[C('Причина отказа')]), 'rejectComment': '',
         'appStatus': None, 'rowCreated': iso_date(r[C('Дата создания строки')], i, 'rowCreated'),
         'source': '', 'sourceLink': s(r[11]), 'unread': False, 'risk': False, 'fav': False,
         'law': None, 'procType': s(r[5]) or None, 'checkoVerified': False,
         'aiReport': None, 'aiReportUrl': s(r[12]) or None,
         'description': '', 'aiVerdict': None, 'aiComment': '', 'documents': []}
    if ar:
        t['submissionType'] = 'На почту' if 'почт' in s(ar[5]).lower() else ('На площадке' if s(ar[5]) else '')
        t['platformLink'] = s(ar[6])
        t['submittedAt'] = iso_date(ar[3], i, 'answer_date')
        t['responseDueDate'] = iso_date(ar[7], i, 'answer_due')
        t['answerComment'] = s(ar[9])
    first_name = 'ПКО' if (st_raw == 'ПКО/НДА' or i in FIRST_IS_PKO) else 'Основной'
    status = None
    if dec == 'Отказ':
        status = {'Отклонена':'Отклонено','не успели податься':'Отказались сами','Отказались сами':'Отказались сами'}.get(st_raw)
        if st_raw == 'не успели податься': t['rejectReason'] = 'Не успели податься'
        if i in PKO_REFUSED:
            t['multiStage'] = True
            t['tenderSections'] = [section(0,'ПКО','Допущены',t['deadline'],t['deadline'][:10],True),
                                   section(1,'Основной','Отклонено','','',True)]
            status = 'Отклонено'
        if st_raw and st_raw not in ('Отклонена','не успели податься','Отказались сами','ПКО/НДА'):
            log['refusal_with_status'].append((i, st_raw))
        if not t['rejectReason']: t['rejectReason'] = 'Прочее'
    elif dec == 'На рассмотрении':
        status = 'На рассмотрении'
        t['rejectReason'] = ''
    elif dec == 'Участвуем':
        status = TENDERS_STATUS.get(st_raw)
        if st_raw and status is None: log['unknown_status'].append((i, st_raw))
        if res == 'прочее':
            status = PROCHEE[i]
        elif res:
            if res not in ANSWER_STATUS: log['unknown_answer_result'].append((i, res))
            else: status = ANSWER_STATUS[res]
        if i == 472:
            status = 'Отказались сами'; t['rejectReason'] = 'Невыполнимые условия'
        if st_raw == 'не успели податься': t['rejectReason'] = 'Не успели податься'
        if status in LOSS:
            if not t['rejectReason']: t['rejectReason'] = 'Прочее'
        else:
            t['rejectReason'] = ''
    else:
        log['unexpected_decision'].append((i, dec))
    t['appStatus'] = status
    # submission dates / stages
    submitted = (dec == 'Участвуем' or i in SUBMITTED_REFUSALS) and status not in NOT_SUBMITTED_STATUSES and st_raw != 'не успели податься' \
        and not (st_raw == 'Отменено' and not ar)
    first_actual = ''
    if dec == 'На рассмотрении' and st_raw == 'ПКО/НДА' and t['deadline']:
        first_actual = t['deadline'][:10]  # дата ПКО из «Тендеров», решение Оксаны
    if submitted:
        first_actual = SUBMIT_DATE_FROM_STATS.get(i) or t.get('submittedAt') or (t['deadline'][:10] if t['deadline'] else '')
        if not first_actual: log['submitted_without_date'].append(i)
    if 'tenderSections' not in t:
        stages = extra.get(i, [])
        if stages:
            t['multiStage'] = True
            secs = [section(0, first_name, 'Допущены', t['deadline'], first_actual, True)]
            for k, (nm, dt) in enumerate(stages):
                last = k == len(stages) - 1
                secs.append(section(k+1, nm, status if last else 'Допущены', '', dt, (status in FINAL) if last else True))
            t['tenderSections'] = secs
        else:
            t['tenderSections'] = [section(0, first_name, status or '', t['deadline'], first_actual, status in FINAL)]
    if i in PRE_STAGE:
        nm, dt = PRE_STAGE[i]
        t['multiStage'] = True
        first = t['tenderSections'][0]
        if first['status'] != 'Допущены' and len(t['tenderSections']) == 1:
            pass
        t['tenderSections'].insert(0, section(0, nm, 'Допущены', t['deadline'], dt, True))
    if i == 1563 and not t['tenderSections'][0].get('actualSubmissionDate'):
        log['pko_without_date'].append(i)
    for k, sec in enumerate(t['tenderSections']): sec['id'] = 'section-%d-%d' % (nid, k+1)
    if status in FINAL or dec == 'Отказ':
        t['archivedAt'] = t['submittedAt'] if t.get('submittedAt') else t['rowCreated']
    if status == 'Выиграли':
        if i in WIN_DATES:
            t['resultDate'] = WIN_DATES[i]; t['archivedAt'] = WIN_DATES[i]
        else:
            log['win_without_date'].append(i)
    out.append(t)

# ---------- new tenders from «Ответы» ----------
for ai in NEW_FROM_ANSWERS:
    ar = answer_rows[ai]; nid += 1
    d = iso_date(ar[3], 'O%d' % ai, 'answer_date')
    t = {'id': nid, 'sheetRow': 'Ответы %d' % ai, 'manager': manager(ar[0]), 'name': s(ar[1]), 'customer': s(ar[2]),
         'nmc': None, 'currencies': ['₽'], 'pub': '', 'deadline': d, 'dept': '', 'depts': [],
         'decision': 'Участвуем', 'rejectReason': '', 'rejectComment': '', 'appStatus': 'Подано',
         'rowCreated': d, 'source': '', 'sourceLink': s(ar[4]), 'unread': False, 'risk': False, 'fav': False,
         'law': None, 'procType': None, 'checkoVerified': False, 'aiReport': None, 'aiReportUrl': None,
         'description': '', 'aiVerdict': None, 'aiComment': '', 'documents': [],
         'submissionType': 'На почту' if 'почт' in s(ar[5]).lower() else 'На площадке',
         'platformLink': s(ar[6]), 'submittedAt': d, 'responseDueDate': iso_date(ar[7], 'O%d' % ai, 'answer_due'),
         'answerComment': s(ar[9]),
         'tenderSections': [{'id': 'section-%d-1' % nid, 'name': 'Основной', 'status': 'Подано', 'deadline': d,
                             'actualSubmissionDate': d, 'completed': False}]}
    out.append(t)

for t in out:
    t['lifecycleSection'] = None
    for k in [k for k, v in t.items() if v in ('', None) and k not in ('manager','dept','decision','rejectReason','rejectComment','appStatus','nmc','law','procType','aiReport','aiReportUrl','aiVerdict','source','sourceLink','description','aiComment','pub','deadline','rowCreated')]:
        del t[k]

js = ("// Реальные тендеры из Google-таблицы «Тендеры» (листы «Тендеры», «Ответы», «статистика»).\n"
      "// Сформировано автоматически " + datetime.date.today().isoformat() + ". Дубли не загружены.\n"
      "// Меняя данные в этом файле, увеличьте TENDER_SEED_VERSION — иначе браузер оставит старую копию.\n"
      "window.TENDER_SEED_VERSION = 'real-2026-09-17-v5';\n"
      "window.TENDER_SEED_DATA = " + json.dumps(out, ensure_ascii=False, separators=(',', ':')) + ";\n")
open(OUT, 'w').write(js)

# ---------- report ----------
cnt = collections.Counter((t['decision'], t['appStatus']) for t in out)
def kind(k, sct):
    if sct['name'] in ('ПКО', 'НДА'): return 'pko'
    if sct['name'] == 'Переторжка': return 'retender'
    return 'first' if k == 0 else 'main'
kinds = collections.Counter(kind(k, sct) for t in out for k, sct in enumerate(t['tenderSections']) if sct.get('actualSubmissionDate'))
first, ret, main = kinds['first'], kinds['retender'], kinds['main']
rep = {'total': len(out), 'excluded_duplicates': len(excluded), 'answers_matched': len(answers),
       'by_decision_status': {'%s / %s' % k: v for k, v in sorted(cnt.items(), key=lambda x: -x[1])},
       'submissions_first': first, 'retenders': ret, 'main_stage': main, 'pko': kinds['pko'],
       'ai_links': sum(1 for t in out if t.get('aiReportUrl')),
       'managers': collections.Counter(t['manager'] for t in out),
       'depts': collections.Counter(t['dept'] for t in out),
       'reasons': collections.Counter(t['rejectReason'] for t in out),
       'log': log, 'js_chars': len(js)}
open(REPORT, 'w').write(json.dumps(rep, ensure_ascii=False, indent=1, default=str))
print(json.dumps({k: v for k, v in rep.items() if k != 'log'}, ensure_ascii=False, indent=1, default=str))
print('LOG', json.dumps(log, ensure_ascii=False, default=str)[:3000])
